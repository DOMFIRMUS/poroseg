"""Stage 'report': assemble everything the author needs to write the paper.

Outputs
-------
FIGURES_AND_TABLES.md   catalogue of every figure/table (main vs supplementary),
                        with draft captions and file names
RESULTS_DIGEST.md       every key number registered by the stages, grouped,
                        with its description (numbers to quote in the text)
key_numbers.json        the same, machine readable
tables_all.xlsx         every table as a worksheet (+ an index sheet)
tables/method_scorecard one row per method combining accuracy, robustness
                        and stability indicators produced by other stages
"""
from __future__ import annotations

import json
import re
import time

import numpy as np
import pandas as pd

from .. import __version__


def _load_json_dir(d):
    out = {}
    for f in sorted(d.glob("*.json")):
        try:
            out[f.stem] = json.loads(f.read_text(encoding="utf-8"))
        except Exception:
            out[f.stem] = None
    return out


def _fmt_value(v, unit=""):
    if isinstance(v, float):
        s = f"{v:.4g}" if abs(v) < 1e4 else f"{v:.3e}"
    elif isinstance(v, (list, tuple)):
        s = "[" + ", ".join(_fmt_value(x) for x in v) + "]"
    elif isinstance(v, dict):
        s = "; ".join(f"{k}: {_fmt_value(x)}" for k, x in v.items())
    else:
        s = str(v)
    return f"{s} {unit}".strip()


STAGE_ORDER = ["segmentation", "descriptive", "spatial", "sensitivity", "perturbation", "morphology",
               "synthetic", "ml_baseline", "inference", "legacy", "report"]


def _scorecard(ctx) -> pd.DataFrame | None:
    """Combine per-method indicators from the other stages (each optional)."""
    rows = {m: {"Method": m} for m in ctx.methods}
    seg = ctx.seg_primary()
    base = seg[seg["post"] == "none"]
    # accuracy vs laboratory (n = 7)
    ref = ctx.lab_ref.set_index("sample_id")["ref"]
    acc = base[base["sample_id"].isin(ref.index)].copy()
    if len(acc):
        acc["err"] = 100 * (acc["porosity"] - acc["sample_id"].map(ref))
        for m, g in acc.groupby("method"):
            rows[m]["Bias vs lab (pp)"] = g["err"].mean()
            rows[m]["MAE vs lab (pp)"] = g["err"].abs().mean()
            rows[m]["n (lab)"] = len(g)
    # inter-sample failures
    for m, g in base.groupby("method"):
        rows[m]["Failures (φ<1% or >90%)"] = int(((g["porosity"] < 0.01) | (g["porosity"] > 0.90)).sum())
    # threshold robustness
    try:
        psi = ctx.read_data("psi")
        for m, g in psi.groupby("method"):
            rows[m]["PSI (median)"] = g["psi"].median()
    except Exception:
        pass
    # spatial consistency
    try:
        sp = ctx.read_data("spatial_ci")
        sp = sp[(sp["roi"] == ctx.primary_roi) & (sp["post"] == "none")]
        for m, g in sp.groupby("method"):
            rows[m]["Spatial CV (median)"] = g["cv_tiles"].median()
    except Exception:
        pass
    # artefact robustness
    try:
        pe = ctx.read_data("perturbation")
        for art in sorted(pe["perturbation"].unique()):
            sub = pe[pe["perturbation"] == art]
            lvl = sorted(sub["level"].unique())
            ref_lvl = lvl[len(lvl) // 2] if len(lvl) > 1 else lvl[0]
            s2 = sub[sub["level"] == ref_lvl]
            for m, g in s2.groupby("method"):
                rows[m][f"|Δφ| {art}={ref_lvl:g} (pp)"] = g["delta_pp"].abs().median()
    except Exception:
        pass
    # synthetic ground truth
    try:
        sy = ctx.read_data("synthetic_benchmark")
        col = next((c for c in ("abs_error_pp", "abs_err_pp", "abs_error") if c in sy.columns), None)
        if col is not None:
            post_col = "post" if "post" in sy.columns else None
            s2 = sy[sy[post_col] == "none"] if post_col else sy
            for m, g in s2.groupby("method"):
                rows[m]["Synthetic |error| (pp, median)"] = g[col].median()
            if "dice" in sy.columns:
                for m, g in s2.groupby("method"):
                    rows[m]["Synthetic Dice (median)"] = g["dice"].median()
    except Exception:
        pass
    df = pd.DataFrame(rows.values())
    return df if df.shape[1] > 1 else None


def _write_xlsx(ctx, catalog_tables):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    idx = wb.active
    idx.title = "index"
    idx.append(["sheet", "table", "kind", "stage", "caption"])
    used = set()
    for t in catalog_tables:
        f = ctx.dirs["tables"] / f"{t['name']}.csv"
        if not f.exists():
            continue
        df = pd.read_csv(f)
        sheet = re.sub(r"[\[\]:*?/\\]", "_", t["name"])[:31]
        k = 1
        while sheet in used:
            sheet = f"{sheet[:28]}_{k}"
            k += 1
        used.add(sheet)
        ws = wb.create_sheet(sheet)
        ws.append(list(map(str, df.columns)))
        for row in df.itertuples(index=False):
            ws.append([None if (isinstance(v, float) and not np.isfinite(v)) else v for v in row])
        for c in ws[1]:
            c.font = Font(bold=True)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for col in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(8, width + 2), 60)
        ws.freeze_panes = "A2"
        idx.append([sheet, t["name"], t.get("kind", ""), t.get("stage", ""), t.get("caption", "")])
    for c in idx[1]:
        c.font = Font(bold=True)
    idx.column_dimensions["E"].width = 120
    wb.save(ctx.out / "tables_all.xlsx")


def run(ctx) -> None:
    sc = _scorecard(ctx)
    if sc is not None:
        ctx.save_table(sc, "method_scorecard",
                       "Per-method summary of accuracy against the laboratory reference (n = 7 images, post-processing "
                       "none, disc ROI), gross failures over the 9 images, threshold stability (PSI), spatial "
                       "consistency (median CV of tile porosity), robustness to controlled artefacts (median absolute "
                       "porosity change at the middle perturbation level) and accuracy on synthetic images with known "
                       "ground truth. Stability indicators do not imply accuracy.", kind="main",
                       float_format="{:.3g}")
    ctx.finalize()   # make sure the scorecard itself is in the catalogue before assembling it

    catalogs = _load_json_dir(ctx.dirs["catalog"])
    keys = _load_json_dir(ctx.dirs["keys"])
    order = [s for s in STAGE_ORDER if s in catalogs] + [s for s in catalogs if s not in STAGE_ORDER]

    # ---------------------------------------------------------- catalogue
    items = [it for s in order for it in (catalogs[s] or [])]
    lines = [f"# Figures and tables (generated {time.strftime('%Y-%m-%d %H:%M')}, poroseg {__version__})", "",
             "Draft captions are generated by the code; edit them in the manuscript, not here.", ""]
    for kind, title in (("main", "Main text"), ("supp", "Supplementary material")):
        for typ in ("figure", "table"):
            sel = [it for it in items if it["type"] == typ and it.get("kind", "main") == kind]
            if not sel:
                continue
            lines.append(f"## {title} - {typ}s ({len(sel)})")
            lines.append("")
            for it in sel:
                folder = "figures" if typ == "figure" else "tables"
                files = ", ".join(f"[{f}]({folder}/{f})" for f in it.get("files", []))
                lines.append(f"### `{it['name']}`  *(stage: {it.get('stage', '')})*")
                lines.append("")
                lines.append(it.get("caption", ""))
                if it.get("notes"):
                    lines.append("")
                    lines.append(f"*Notes:* {it['notes']}")
                lines.append("")
                lines.append(f"Files: {files}")
                lines.append("")
    (ctx.out / "FIGURES_AND_TABLES.md").write_text("\n".join(lines), encoding="utf-8")

    # --------------------------------------------------------- key numbers
    merged = {}
    dl = [f"# Results digest (generated {time.strftime('%Y-%m-%d %H:%M')})", "",
          "Every number below was produced by the pipeline and is traceable to the stage that computed it. "
          "Quote these values (not hand-typed ones) in the manuscript.", ""]
    for s in [s for s in STAGE_ORDER if s in keys] + [s for s in keys if s not in STAGE_ORDER]:
        if not keys[s]:
            continue
        dl.append(f"## {s}")
        dl.append("")
        dl.append("| key | value | description |")
        dl.append("|---|---|---|")
        for k, v in keys[s].items():
            merged[k] = v
            desc = str(v.get("desc", "")).replace("|", "/")
            dl.append(f"| `{k}` | {_fmt_value(v.get('value'), v.get('unit', ''))} | {desc} |")
        dl.append("")
    (ctx.out / "RESULTS_DIGEST.md").write_text("\n".join(dl), encoding="utf-8")
    (ctx.out / "key_numbers.json").write_text(json.dumps(merged, indent=1, ensure_ascii=False), encoding="utf-8")

    _write_xlsx(ctx, [it for it in items if it["type"] == "table"])
    n_fig = sum(it["type"] == "figure" for it in items)
    n_tab = sum(it["type"] == "table" for it in items)
    ctx.log("catalogue: %d figures, %d tables, %d key numbers", n_fig, n_tab, len(merged))
